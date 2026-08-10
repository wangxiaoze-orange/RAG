"""文档解析 v2：新增 HTML / Excel / CSV 解析支持
- 独立模块，不影响原 parser.py 的稳定性
- parser.py 通过 try_parse_v2 委托调用
"""
import io
import logging

logger = logging.getLogger(__name__)


def _decode_text(data: bytes) -> str:
    """兼容 UTF-8 / GBK / GB18030 编码"""
    for enc in ("utf-8", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_html(data: bytes) -> str:
    from bs4 import BeautifulSoup
    html_text = _decode_text(data)
    soup = BeautifulSoup(html_text, "lxml")
    # 移除 script / style / nav / footer 等非正文标签
    for tag in soup.find_all(["script", "style", "nav", "footer", "header", "noscript", "iframe"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)


def _parse_excel(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    max_rows = 50000
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if sheet_name != wb.sheetnames[0]:
            parts.append(f"\n## {sheet_name}\n")
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if row_count >= max_rows:
                logger.warning("Excel 行数超过 %d，已截断", max_rows)
                break
            cells = [str(c).strip() if c is not None else "" for c in row]
            if any(cells):  # 跳过全空行
                parts.append(" | ".join(cells))
            row_count += 1
    wb.close()
    return "\n".join(parts)


def _parse_csv(data: bytes) -> str:
    import csv
    text = _decode_text(data)
    reader = csv.reader(text.splitlines())
    parts = []
    for row in reader:
        if any(cell.strip() for cell in row):  # 跳过全空行
            parts.append(" | ".join(cell.strip() for cell in row))
    return "\n".join(parts)


def try_parse_v2(ext: str, data: bytes) -> tuple[str | None, str]:
    """尝试用 v2 新增解析器处理，返回 (text, parser_name)；不支持的扩展返回 (None, '')"""
    if ext in ("html", "htm"):
        try:
            text = _parse_html(data)
            return text, "html"
        except Exception as e:  # noqa: BLE001
            logger.warning("html 解析失败，按文本处理: %s", e)
            return _decode_text(data), "text"
    elif ext in ("xlsx", "xls"):
        return _parse_excel(data), "excel"
    elif ext == "csv":
        return _parse_csv(data), "csv"
    return None, ""
